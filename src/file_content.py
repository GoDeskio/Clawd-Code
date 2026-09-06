"""Safe text extraction for user-selected local files.

The desktop stages a private copy before calling these helpers.  Extraction is
bounded so an uploaded document cannot silently consume the model context.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
import zipfile


RICH_DOCUMENT_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".zip"}


def _bounded(parts: list[str], max_chars: int) -> tuple[str, bool]:
    text = "\n".join(part for part in parts if str(part).strip())
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars].rstrip() + "\n[content truncated]", True


def extract_rich_text(path: str | Path, *, max_chars: int = 256_000) -> tuple[bool, str, str]:
    """Return ``(supported, text, note)`` for common document containers."""
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix not in RICH_DOCUMENT_SUFFIXES:
        return False, "", ""
    try:
        parts: list[str] = []
        kind = suffix.lstrip(".").upper()
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(target))
            for index, page in enumerate(reader.pages, 1):
                parts.append(f"--- Page {index} ---\n{page.extract_text() or ''}")
        elif suffix == ".docx":
            from docx import Document

            document = Document(str(target))
            parts.extend(paragraph.text for paragraph in document.paragraphs)
            for table_index, table in enumerate(document.tables, 1):
                parts.append(f"--- Table {table_index} ---")
                parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        elif suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(str(target), read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    parts.append(f"--- Sheet: {sheet.title} ---")
                    for row in sheet.iter_rows(values_only=True):
                        values = ["" if value is None else str(value) for value in row]
                        if any(values):
                            parts.append("\t".join(values))
                        if sum(len(item) for item in parts) >= max_chars:
                            break
                    if sum(len(item) for item in parts) >= max_chars:
                        break
            finally:
                workbook.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(str(target))
            for index, slide in enumerate(presentation.slides, 1):
                parts.append(f"--- Slide {index} ---")
                for shape in slide.shapes:
                    text = getattr(shape, "text", "")
                    if text:
                        parts.append(str(text))
        elif suffix == ".zip":
            with zipfile.ZipFile(target) as archive:
                parts = [f"{item.filename}\t{item.file_size} bytes" for item in archive.infolist()]
                kind = "ZIP file listing"
        text, truncated = _bounded(parts, max(1, int(max_chars)))
        note = f"extracted {kind} content"
        if truncated:
            note += f" (truncated to {max_chars} characters)"
        if not text.strip():
            note = f"{kind} contained no extractable text"
        return True, text, note
    except Exception as exc:  # malformed documents remain downloadable/tool-readable
        return True, "", f"could not extract {suffix.lstrip('.').upper()} content: {exc}"
