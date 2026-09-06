"""Jonathan-native raster-to-editable SVG and PPTX reconstruction."""

from __future__ import annotations

import html
import importlib.util
import json
from pathlib import Path
from typing import Any, Callable

DRAWAI_REPOSITORY = "https://github.com/Renaissance-Mind/DrawAI.git"
DRAWAI_LICENSE = "Apache-2.0"


class DrawAiManager:
    """Compatibility name retained for a standalone Jonathan implementation."""

    def __init__(self, source_dir: str | Path) -> None:
        self.source_dir = Path(source_dir).expanduser().resolve()
        self.checkout = self.source_dir / "src" / "integrations"
        self.runtime = self.source_dir / "src"

    @staticmethod
    def _dependencies() -> tuple[bool,list[str]]:
        missing=[module for module in ("PIL","cv2","numpy","pptx") if importlib.util.find_spec(module) is None]
        return not missing,missing

    def status(self) -> dict[str, Any]:
        ready,missing=self._dependencies()
        return {"available": True,"source_ready": True,"dependencies_ready": ready,"models_ready": True,"runtime_ready": ready,
                "path":str(self.checkout),"runtime_path":str(self.runtime),"repository":"internal://jonathan/editable-graphics",
                "reference_repository":DRAWAI_REPOSITORY,"revision":"jonathan-native-1","license":"MIT (implementation); Apache-2.0 design reference",
                "payments_required":False,"model_terms_apply":False,"disk_requirement":"No separate model pack", "missing":missing,
                "features":["color-region vectorization","shape contours","optional OCR","editable SVG","native-shape PPTX"]}

    def sync(self, progress: Callable[[str],None]|None=None) -> dict[str,Any]:
        if progress: progress("Jonathan's native editable-graphics engine is ready")
        return self.status()

    def install(self, *, models: bool=False, device: str="cpu", progress: Callable[[str],None]|None=None) -> dict[str,Any]:
        status=self.sync(progress)
        if not status["runtime_ready"]: raise RuntimeError(f"Missing Jonathan image/document dependencies: {', '.join(status['missing'])}")
        return status

    @staticmethod
    def _ocr(image_path: Path) -> list[dict[str,Any]]:
        try:
            from rapidocr_onnxruntime import RapidOCR
            result,_=RapidOCR()(str(image_path))
            rows=[]
            for box,text,score in result or []:
                xs=[float(p[0]) for p in box]; ys=[float(p[1]) for p in box]
                rows.append({"text":str(text),"score":float(score),"x":min(xs),"y":min(ys),"width":max(xs)-min(xs),"height":max(ys)-min(ys)})
            return rows
        except Exception: return []

    def convert(self, image: str|Path, output_root: str|Path, *, device: str="cpu", timeout: int=7200) -> dict[str,Any]:
        self.install(); source=Path(image).expanduser().resolve(); target=Path(output_root).expanduser().resolve()
        if not source.is_file(): raise ValueError(f"Input image does not exist: {source}")
        target.mkdir(parents=True,exist_ok=True)
        from PIL import Image,ImageDraw
        import cv2
        import numpy as np
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches,Pt

        with Image.open(source) as opened: rgb=opened.convert("RGB")
        width,height=rgb.size; scale=min(1.0,1200/max(width,height)); work=rgb.resize((max(1,int(width*scale)),max(1,int(height*scale))))
        array=np.array(work); small=work.quantize(colors=10,method=Image.Quantize.MEDIANCUT).convert("RGB"); palette=np.array(small)
        colors,counts=np.unique(palette.reshape(-1,3),axis=0,return_counts=True); order=np.argsort(counts)[::-1]
        svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',f'<rect width="{width}" height="{height}" fill="rgb({colors[order[0]][0]},{colors[order[0]][1]},{colors[order[0]][2]})"/>']
        regions=[]; inv=1/scale
        for color in colors[order[1:]]:
            mask=np.all(palette==color,axis=2).astype(np.uint8)*255
            contours,_=cv2.findContours(mask,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)
            for contour in contours:
                if cv2.contourArea(contour)<40: continue
                approx=cv2.approxPolyDP(contour,1.5,True).reshape(-1,2)
                points=[(round(float(x)*inv,1),round(float(y)*inv,1)) for x,y in approx]
                if len(points)<3: continue
                fill=f"rgb({int(color[0])},{int(color[1])},{int(color[2])})"; svg.append(f'<polygon points="{" ".join(f"{x},{y}" for x,y in points)}" fill="{fill}"/>')
                x,y,w,h=cv2.boundingRect(contour); regions.append({"x":x*inv,"y":y*inv,"width":w*inv,"height":h*inv,"color":[int(v) for v in color]})
        ocr=self._ocr(source)
        for row in ocr:
            size=max(8,min(72,row["height"]*.78)); svg.append(f'<text x="{row["x"]}" y="{row["y"]+row["height"]}" font-family="Arial, sans-serif" font-size="{size}" fill="#111">{html.escape(row["text"])}</text>')
        svg.append("</svg>"); svg_path=target/f"{source.stem}-editable.svg"; svg_path.write_text("\n".join(svg),encoding="utf-8")

        prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(13.333*height/width); slide=prs.slides.add_slide(prs.slide_layouts[6])
        sx=prs.slide_width/width; sy=prs.slide_height/height
        background=colors[order[0]]; shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height); shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor(*[int(v) for v in background]); shape.line.fill.background()
        for region in sorted(regions,key=lambda r:r["width"]*r["height"],reverse=True)[:400]:
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,int(region["x"]*sx),int(region["y"]*sy),max(1,int(region["width"]*sx)),max(1,int(region["height"]*sy))); shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor(*region["color"]); shape.line.fill.background()
        for row in ocr:
            box=slide.shapes.add_textbox(int(row["x"]*sx),int(row["y"]*sy),max(1,int(row["width"]*sx)),max(1,int(row["height"]*sy*1.25))); paragraph=box.text_frame.paragraphs[0]; paragraph.text=row["text"]; paragraph.font.name="Arial"; paragraph.font.size=Pt(max(6,row["height"]*0.56))
        pptx_path=target/f"{source.stem}-editable.pptx"; prs.save(pptx_path)
        preview=rgb.copy(); draw=ImageDraw.Draw(preview)
        for row in ocr: draw.rectangle((row["x"],row["y"],row["x"]+row["width"],row["y"]+row["height"]),outline="#e14949",width=2)
        preview_path=target/f"{source.stem}-review.png"; preview.save(preview_path)
        package_path=target/"editable_graphics_package.json"; package_path.write_text(json.dumps({"source":str(source),"size":[width,height],"regions":regions,"ocr":ocr,"outputs":[str(svg_path),str(pptx_path),str(preview_path)]},ensure_ascii=False,indent=2),encoding="utf-8")
        return {**self.status(),"input":str(source),"run_dir":str(target),"output":f"Vectorized {len(regions)} color regions and {len(ocr)} text regions.","artifacts":[str(svg_path),str(pptx_path),str(preview_path),str(package_path)]}
